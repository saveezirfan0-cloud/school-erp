#!/usr/bin/env python3
"""Extract the monthly income / expense ledger from the school's Excel
accounts workbook ("<name>_Accounts.xlsx") into compact JSON entries that
can be loaded into the `payments` / `expenses` tables.

Usage:
    python3 scripts/extract_workbook_ledger.py path/to/Accounts.xlsx > ledger.json

Each entry is [date, kind, category, branch, amount, summary]:
    kind     "i" (income / cash in) or "e" (expense / cash out)
    branch   "B" Baneen, "N" Bnat, "S" School, "" none (org-level / Umer Colony)
    summary  1 when only a monthly total exists (January), else 0

Sources
    Jan-2026 Statement      monthly totals only (dated month end)
    Feb..Sep-2026 sheets    one entry per category per day

Student fee income for August onwards ("Baneen Fees") is NOT exported:
those collections already live in the ERP as invoice payments, so
importing them would double count. Everything else is exported.
"""
import datetime
import json
import sys

import openpyxl

MONTH_SHEETS = [
    "Feb-2026", "Mar-2026", "Apr-2026", "May-2026", "June-2026",
    "July-2026", "August-2026", "September-2026",
]
# First month whose student fee collections are tracked in the ERP itself.
ERP_FEES_FROM = datetime.date(2026, 8, 1)

RENAME = {
    "baneen fes": "Baneen Fees", "baneen fees": "Baneen Fees",
    "welfear": "Welfare", "welfare": "Welfare",
    "banaat addmission fees": "Banaat Admission Fees",
    "baneen staf salary": "Baneen Staff Salary",
    "baneen staff salary": "Baneen Staff Salary",
    "mantenance and repairing": "Maintenance and Repairing",
    "furniature + fixture": "Furniture + Fixture",
    "baneen stanationary": "Baneen Stationary",
    "utility baneen": "Utility Bill Baneen",
    "tution fee": "Tuition Fee",
    "baneen printing": "Baneen Printing & Designing",
    "pt,cash in hand": "PT Cash in Hand", "pt,cash": "PT Cash", "pt cash": "PT Cash",
    "suspense": "Suspense", "lunch exp": "Lunch Expense",
}


def clean(label):
    return RENAME.get(label.strip().lower(), label.strip())


def branch_of(cat):
    c = cat.lower()
    if "banaat" in c or "bnat" in c:
        return "N"
    if "school" in c:
        return "S"
    if "umer" in c:
        return ""
    if c in ("welfare", "donation", "loan", "loan return", "suspense", "miscellaneous",
             "pt cash", "pt cash in hand", "baneen adv"):
        return ""
    return "B"  # everything else sits in the Baneen block of the sheet


def month_sheet_entries(ws):
    hdr = next(r for r in range(1, 12)
               if sum(isinstance(c.value, datetime.datetime) for c in ws[r]) > 10)
    dates = {c.column: c.value.date() for c in ws[hdr] if isinstance(c.value, datetime.datetime)}
    kind = "i"
    for r in range(hdr + 1, ws.max_row + 1):
        label = ws.cell(r, 1).value
        if ws.cell(r, 4).value and str(ws.cell(r, 4).value).strip().lower().startswith("expens"):
            kind = "e"
        if not isinstance(label, str) or not label.strip():
            continue
        low = label.strip().lower()
        if low.startswith("total"):
            continue
        if low in ("expense", "expenses"):
            kind = "e"
            continue
        cat = clean(label)
        for col, d in dates.items():
            v = ws.cell(r, col).value
            if isinstance(v, (int, float)) and v:
                if kind == "i" and cat == "Baneen Fees" and d >= ERP_FEES_FROM:
                    continue
                yield [d.isoformat(), kind, cat, branch_of(cat), v, 0]


def january_entries(ws):
    d = "2026-01-31"
    for r in range(7, 17):                       # income block (A/B)
        label, v = ws.cell(r, 1).value, ws.cell(r, 2).value
        if isinstance(label, str) and isinstance(v, (int, float)) and v:
            cat = "Baneen Fees" if label.strip() == "Fees" else clean(label)
            yield [d, "i", cat, branch_of(cat), v, 1]
    for r in range(21, 40):                      # expense block (A/B)
        label, v = ws.cell(r, 1).value, ws.cell(r, 2).value
        if isinstance(label, str) and isinstance(v, (int, float)) and v:
            cat = "Baneen Staff Salary" if label.strip() == "Salaries" else clean(label)
            yield [d, "e", cat, branch_of(cat), v, 1]


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    out = list(january_entries(wb["Jan-2026 Statement"]))
    for name in MONTH_SHEETS:
        out.extend(month_sheet_entries(wb[name]))
    json.dump(out, sys.stdout, separators=(",", ":"))


if __name__ == "__main__":
    main(sys.argv[1])
